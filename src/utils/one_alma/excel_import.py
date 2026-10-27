import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from uuid import UUID, uuid4
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from ..excel_validation import validate_xlsx_upload, validation_message
from .models import NewProduct, ShipmentItem, shipment_fingerprint

IMPORT_HEADERS = ("Артикул", "Наименование", "Объём", "Годен до", "Цена ЦБ", "Цена опт", "Цена магазина", "Количество")


@dataclass(frozen=True)
class Shipment:
    products: list[NewProduct]
    rows: list[dict[str, object]]
    fingerprint: str
    receipt_id: str

    @property
    def valid(self) -> bool:
        return bool(self.products) and all(not row["error"] for row in self.rows)


class ExcelProductImporter:
    max_file_size = 5 * 1024 * 1024
    max_rows = 2000

    @staticmethod
    def template() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Импорт"
        sheet.append(IMPORT_HEADERS)
        sheet.freeze_panes = "A2"
        for column in "ABCDEFGH":
            sheet.column_dimensions[column].width = 22 if column != "B" else 55
        sheet.column_dimensions["A"].number_format = "@"
        metadata = workbook.create_sheet("_OneAlma")
        metadata.append(["shipment_id", str(uuid4())])
        metadata.sheet_state = "hidden"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def parse(self, content: bytes, filename: str) -> list[NewProduct]:
        shipment = self.inspect(content, filename)
        self.require_valid(shipment)
        return shipment.products

    @staticmethod
    def require_valid(shipment: Shipment) -> None:
        errors = [f"Строка {row['row']}: {row['error']}" for row in shipment.rows if row["error"]]
        if errors:
            raise ValueError("\n".join(errors) + " Импорт отменён.")

    def inspect(self, content: bytes, filename: str) -> Shipment:
        validate_xlsx_upload(content, filename, self.max_file_size)
        self._validate_archive(content)
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except Exception:
            raise ValueError("Не удалось прочитать Excel. Проверьте файл .xlsx.") from None
        try:
            sheet = workbook.worksheets[0]
            if (sheet.max_row or 0) > self.max_rows + 20 or (sheet.max_column or 0) > 30:
                raise ValueError("В файле допускается до 2000 строк товаров и 30 колонок.")
            headers = None
            products: list[NewProduct] = []
            rows: list[dict[str, object]] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if headers is None:
                    candidate = self._headers(row)
                    if candidate:
                        headers = candidate
                    elif row_number > 10:
                        break
                    continue
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                if str(row[0]).upper().startswith("ИТОГО"):
                    continue
                values = {field: row[column] for field, column in headers.items()}
                view: dict[str, object] = {field: self._text(value)[:300] for field, value in values.items()}
                view.update(row=row_number, error="")
                try:
                    labels = dict(zip(ShipmentItem.model_fields, IMPORT_HEADERS))
                    formula_fields = [
                        field for field, value in values.items() if isinstance(value, str) and value.startswith("=")
                    ]
                    if formula_fields:
                        raise ValueError("; ".join(f"{labels[field]}: замените формулу значением" for field in formula_fields))
                    item = ShipmentItem(
                        **{field: self._text(value) for field, value in values.items() if field not in ("quantity", "expiry")},
                        expiry=self._expiry(values["expiry"]),
                        quantity=self._integer(values["quantity"]),
                    )
                    products.append(item.to_product())
                    view.update(item.model_dump())
                except ValidationError as exc:
                    view["error"] = validation_message(exc, dict(zip(ShipmentItem.model_fields, IMPORT_HEADERS)))
                except (ValueError, TypeError) as exc:
                    view["error"] = str(exc)
                rows.append(view)
            if not headers:
                raise ValueError("Не найдены заголовки. Скачайте шаблон импорта.")
            if not rows:
                raise ValueError("В файле нет товаров.")
            if len(rows) > self.max_rows or sum(product.quantity for product in products) > 10000:
                raise ValueError("За один импорт допускается до 2000 позиций и 10000 штук.")
            fingerprint = shipment_fingerprint(products)
            receipt_id = "content:" + fingerprint
            if "_OneAlma" in workbook.sheetnames:
                metadata = workbook["_OneAlma"]
                if metadata["A1"].value != "shipment_id":
                    raise ValueError("Повреждён идентификатор поставки. Скачайте новый шаблон.")
                try:
                    receipt_id = str(UUID(str(metadata["B1"].value)))
                except ValueError:
                    raise ValueError("Повреждён идентификатор поставки. Скачайте новый шаблон.") from None
            return Shipment(products, rows, fingerprint, receipt_id)
        finally:
            workbook.close()

    @staticmethod
    def _validate_archive(content: bytes) -> None:
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Excel слишком большой в распакованном виде.")
        except BadZipFile:
            raise ValueError("Некорректный файл .xlsx.") from None

    @staticmethod
    def _headers(row: tuple[object, ...]) -> dict[str, int]:
        aliases = {
            "article": ("арт.", "артикул"),
            "name": ("наименование", "наименованиепрепарата"),
            "volume": ("объем",),
            "expiry": ("годендо", "годендо(окончание)"),
            "quantity": ("количество", "количествошт."),
            "central_bank_price": ("ценацб", "ценапоцбрфна27.05(1шт.,₽)"),
            "wholesale_price": ("ценаопт", "ценаоптemansi(1шт.,₽)"),
            "market_price": ("ценамагазина", "средняяценанарынкевмагазинах(1шт.,₽)"),
        }
        result = {}
        for column, value in enumerate(row):
            name = re.sub(r"\s+", "", str(value or "")).casefold().replace("ё", "е")
            for field, names in aliases.items():
                if name in names:
                    if field in result:
                        raise ValueError("В Excel повторяются колонки.")
                    result[field] = column
        required = set(aliases)
        return result if required <= result.keys() else {}

    @staticmethod
    def _text(value: object) -> str:
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return "" if value is None else str(value).strip()

    @staticmethod
    def _integer(value: object) -> int:
        text = str(value)
        if not re.fullmatch(r"\d+(?:\.0+)?", text):
            raise ValueError("Количество: укажите целое положительное число")
        return int(text.split(".", 1)[0])

    @staticmethod
    def _expiry(value: object) -> str:
        if isinstance(value, date):
            return value.strftime("%Y-%m")
        text = str(value)
        return text
