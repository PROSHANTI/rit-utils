"""Run in a disposable ARM64 rit-utils container with network disabled.

No private templates, application .env, credentials or SMTP are involved.
This checks the real CertificateGenerator + LibreOffice conversion using
an artificial PPTX and leaves /tmp/arm64-certificate.pdf for docker cp.
"""

import asyncio
import importlib.util
import json
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import cv2
import fastapi
import lxml.etree  # type: ignore[import-untyped]
import numpy as np
import PIL
from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.util import Inches, Pt

from src.utils.one_alma.pdf_export import InventoryPdfExporter
from src.utils.one_alma.repository import InventoryRepository

# Load the certificate implementation directly, as in the original smoke test.
spec = importlib.util.spec_from_file_location("arm64_smoke_certificate", "/app/src/utils/gen_cert/gen_cert_handler.py")
assert spec is not None and spec.loader is not None
certificate_module = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = certificate_module
spec.loader.exec_module(certificate_module)
CertificateGenerator = certificate_module.CertificateGenerator
CertificateRequest = certificate_module.CertificateRequest
LibreOfficeConverter = certificate_module.LibreOfficeConverter


def run(*args: str) -> str:
    return subprocess.run(args, check=True, capture_output=True, text=True).stdout.strip()


image = np.zeros((16, 16, 3), dtype=np.uint8)
assert cv2.cvtColor(image, cv2.COLOR_BGR2GRAY).shape == (16, 16)

baseline = {
    "machine": platform.machine(),
    "python": sys.version.split()[0],
    "opencv": cv2.__version__,
    "numpy": np.__version__,
    "pillow": PIL.__version__,
    "lxml": lxml.etree.LXML_VERSION,
    "fastapi": fastapi.__version__,
    "libreoffice": run("libreoffice", "--version"),
    "times_new_roman_fallback": run("fc-match", "-f", "%{family}|%{file}\n", "Times New Roman"),
    "locales": run("locale", "-a").splitlines(),
}
for line in Path("/etc/os-release").read_text().splitlines():
    key, _, value = line.partition("=")
    if key in {"ID", "VERSION_ID", "PRETTY_NAME"}:
        baseline[key] = value.strip('"')

presentation = Presentation()
presentation.slide_width = Inches(10)
presentation.slide_height = Inches(7.5)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])

title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.4), Inches(0.8))
title_run = title.text_frame.paragraphs[0].add_run()
title_run.text = "Проверка сертификата ARM64"
title_run.font.name = "Times New Roman"
title_run.font.size = Pt(28)

box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
box.text_frame.clear()
for i, placeholder in enumerate(("name", "price", "serial")):
    paragraph = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
    paragraph.space_after = Pt(18)
    text_run = paragraph.add_run()
    text_run.text = placeholder
    text_run.font.name = "Times New Roman"
    text_run.font.size = Pt(24)

template = Path("/tmp/arm64-synthetic-template.pptx")
presentation.save(template)


class RecordingConverter:
    def __init__(self) -> None:
        self.substituted_text = ""
        self.converter = LibreOfficeConverter()

    def convert(self, pptx_path: str, pdf_path: str) -> None:
        replaced = Presentation(pptx_path)
        self.substituted_text = "\n".join(shape.text for shape in replaced.slides[0].shapes if shape.has_text_frame)
        self.converter.convert(pptx_path, pdf_path)


converter = RecordingConverter()
response = CertificateGenerator(template_path=template, converter=converter).generate(
    CertificateRequest(name="Тестовый сертификат ARM64", price="1500")
)
assert isinstance(response, FileResponse), "CertificateGenerator failed"
assert "Тестовый сертификат ARM64" in converter.substituted_text
assert "1500 ₽" in converter.substituted_text
assert re.search(r"\b\d{6}\b", converter.substituted_text)
assert not re.search(r"\b(name|price|serial)\b", converter.substituted_text)

pdf = Path(response.path)
assert pdf.read_bytes().startswith(b"%PDF-")
assert pdf.stat().st_size > 2000
output = Path("/tmp/arm64-certificate.pdf")
shutil.copyfile(pdf, output)
baseline["synthetic_pdf_bytes"] = output.stat().st_size
baseline["substitutions_ok"] = True
baseline["private_template_used"] = False
baseline["network_enabled"] = False

with tempfile.TemporaryDirectory(prefix="one-alma-arm64-") as temporary:
    database = Path(temporary) / "one_alma.sqlite3"
    inventory_repository = InventoryRepository(database)
    inventory = inventory_repository.snapshot()
    assert (inventory.total, inventory.sold, inventory.available) == (50, 0, 50)
    inventory = inventory_repository.set_sold(inventory.products[0].units[1].id, True)
    assert (inventory.sold, inventory.available) == (1, 49)
    assert InventoryRepository(database).snapshot().sold == 1
    inventory = inventory_repository.set_showcase_quantity(inventory.products[0].id, 1)
    assert InventoryRepository(database).snapshot().showcase_quantity == 1
    inventory = inventory_repository.set_sold_count(inventory.products[0].id, 2)
    assert (inventory.sold, inventory.available, inventory.showcase_quantity) == (2, 48, 0)
    assert len(inventory.products) == 25
    inventory_pdf = InventoryPdfExporter().export(inventory)
    assert inventory_pdf.startswith(b"%PDF-") and len(inventory_pdf) > 2000
    baseline["one_alma_pdf_bytes"] = len(inventory_pdf)
    baseline["one_alma_sqlite_persistence_ok"] = True

Path("/tmp/arm64-runtime-baseline.json").write_text(json.dumps(baseline, ensure_ascii=False, indent=2) + "\n")
if response.background is not None:
    asyncio.run(response.background())
template.unlink()
print(json.dumps(baseline, ensure_ascii=False, indent=2))
