import re
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from src.main import application
from src.utils.one_alma.excel_import import IMPORT_HEADERS
from src.utils.one_alma.repository import InventoryRepository


@pytest.fixture
def inventory_repository(tmp_path, monkeypatch):
    repository = InventoryRepository(tmp_path / "inventory.sqlite3")
    monkeypatch.setattr(application.routes.one_alma, "repository", repository)
    return repository


@pytest.fixture
def csrf(authenticated_client, inventory_repository):
    response = authenticated_client.get("/one_alma")
    assert response.status_code == 200
    match = re.search(r'data-csrf-token="([^"]+)"', response.text)
    assert match is not None
    return match.group(1)


def product_data(csrf):
    return dict(
        position="26",
        article="0007",
        name="Новый крем",
        volume="50 мл",
        expiry="2028-09",
        central_bank_price="100,25",
        wholesale_price="200",
        market_price="300",
        quantity="3",
        csrf_token=csrf,
    )


def excel_file(*rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        headers or ["№", "Артикул", "Наименование", "Объём", "Годен до", "Цена ЦБ", "Цена опт", "Цена магазина", "Количество"]
    )
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return ("products.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_page_home_card_and_button_placement(authenticated_client, inventory_repository):
    home = authenticated_client.get("/home").text
    response = authenticated_client.get("/one_alma")

    html = response.text
    assert "Учет OneAlma" in home and "Учет продажи товаров" in home
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert html.count('data-count-kind="sold"') == 0
    assert html.count('data-count-kind="showcase"') == 25
    assert html.count('class="icon-button record-sale"') == html.count('class="icon-button product-history"') == 25
    assert "Штука 1" not in html and "data-unit-id=" not in html
    assert '<span class="count-text">0 из 2</span>' in html
    assert 'class="row-available"' not in html
    assert "+ Выставить" not in html and 'id="count-plus"' in html and 'id="count-minus"' in html
    assert "Денежные итоги" not in html and "Поле необязательное" not in html and "из оставшихся штук" not in html
    assert "Стоимость товаров по ценам исходной таблицы." not in html
    assert "Выставлено ·" not in html and "Не выставлено" not in html
    assert "Каждый чекбокс —" not in html and "Изменения сохраняются автоматически." not in html
    assert html.index("Расчёт") < html.index("Скачать PDF") < html.index('class="inventory-table"')
    assert "Импорт товаров" in html and "Удалить выбранные" in html
    assert "Без цены: 4 шт." in html
    assert len(re.findall("data-sort-column=", html)) == 10
    assert 'data-sort-column="5">Цена ЦБ, ₽ <span' in html
    assert "Цена по ЦБ РФ на 27.05, ₽" not in html


@pytest.mark.parametrize("indices, available", [([1], 1), ([1, 0], 0)])
def test_sale_endpoint_counts_individual_units(authenticated_client, inventory_repository, csrf, indices, available):
    units = inventory_repository.snapshot().products[0].units

    responses = [
        authenticated_client.post(f"/one_alma/units/{units[index].id}", data={"sold": "true", "csrf_token": csrf})
        for index in indices
    ]
    page = authenticated_client.get("/one_alma").text

    assert all(response.status_code == 200 for response in responses)
    assert responses[-1].json()["available"] == 48 + available
    assert responses[-1].json()["products"][0]["available"] == available
    assert responses[0].json()["valuations"][0]["amounts"][1] == "1 343,43"
    assert f'<span class="count-text">{len(indices)} из 2</span>' in page


def test_create_endpoint_adds_product_and_renders_fields(authenticated_client, inventory_repository, csrf):
    data = product_data(csrf)

    response = authenticated_client.post("/one_alma/products", data=data)
    read_page = authenticated_client.get("/one_alma")

    assert response.status_code == 201
    assert inventory_repository.snapshot().total == 53
    assert read_page.status_code == 200
    assert "Новый крем" in read_page.text
    assert "0007" in read_page.text


def test_edit_endpoint_updates_existing_product(authenticated_client, inventory_repository, csrf):
    authenticated_client.post("/one_alma/products", data=product_data(csrf))
    added = inventory_repository.snapshot().products[-1]
    data = product_data(csrf) | {"name": "Изменённый крем", "quantity": "4"}

    response = authenticated_client.put(f"/one_alma/products/{added.id}", data=data)

    assert response.status_code == 200
    assert inventory_repository.snapshot().products[-1].cells[2] == "Изменённый крем"
    assert inventory_repository.snapshot().total == 54


def test_bulk_delete_endpoint_removes_selected_products(authenticated_client, inventory_repository, csrf):
    authenticated_client.post("/one_alma/products", data=product_data(csrf))
    added = inventory_repository.snapshot().products[-1]

    response = authenticated_client.request(
        "DELETE", "/one_alma/products", data={"product_ids": [added.id, 1], "csrf_token": csrf}
    )

    assert response.status_code == 200
    assert len(inventory_repository.snapshot().products) == 24
    assert inventory_repository.snapshot().total == 48


def test_delete_one_product_removes_its_sales_and_survives_reopen(authenticated_client, inventory_repository, csrf):
    first = inventory_repository.snapshot().products[0]
    authenticated_client.post(f"/one_alma/units/{first.units[0].id}", data={"sold": "true", "csrf_token": csrf})

    response = authenticated_client.request("DELETE", "/one_alma/products", data={"product_ids": [first.id], "csrf_token": csrf})

    reopened = InventoryRepository(inventory_repository.database).snapshot()
    assert response.status_code == 200
    assert (len(reopened.products), reopened.total, reopened.sold, reopened.available) == (24, 48, 0, 48)
    assert all(product.id != first.id for product in reopened.products)


def test_edit_cannot_reduce_quantity_below_sales(authenticated_client, inventory_repository, csrf):
    first = inventory_repository.snapshot().products[0]
    for unit in first.units:
        authenticated_client.post(f"/one_alma/units/{unit.id}", data={"sold": "true", "csrf_token": csrf})

    response = authenticated_client.put(f"/one_alma/products/{first.id}", data=product_data(csrf) | {"quantity": "1"})

    unchanged = inventory_repository.snapshot().products[0]
    assert response.status_code == 422
    assert unchanged.cells == first.cells
    assert unchanged.sold_count == 2


def test_csrf_rejects_mutations_without_changing_database(authenticated_client, inventory_repository, csrf):
    before = inventory_repository.snapshot()
    file = excel_file([1, "0001", "Крем", "50 мл", "09.2028", 100, 200, 300, 4])

    responses = [
        authenticated_client.post("/one_alma/products/1/sold", data={"quantity": "1", "csrf_token": "wrong"}),
        authenticated_client.post("/one_alma/products/1/showcase", data={"quantity": "1", "csrf_token": "wrong"}),
        authenticated_client.post("/one_alma/units/1", data={"sold": "true", "csrf_token": "wrong"}),
        authenticated_client.post("/one_alma/products", data=product_data("wrong")),
        authenticated_client.request("DELETE", "/one_alma/products", data={"product_ids": [1], "csrf_token": "wrong"}),
        authenticated_client.post("/one_alma/import", data={"csrf_token": "wrong"}, files={"file": file}),
    ]

    assert all(response.status_code == 403 for response in responses)
    assert inventory_repository.snapshot() == before


@pytest.mark.parametrize("changes, status", [({"market_price": "-1"}, 422), ({"quantity": "-1"}, 422)])
def test_duplicate_position_and_invalid_fields_do_not_mutate(authenticated_client, inventory_repository, csrf, changes, status):
    before = inventory_repository.snapshot()
    data = product_data(csrf) | changes

    response = authenticated_client.post("/one_alma/products", data=data)

    assert response.status_code == status
    assert inventory_repository.snapshot() == before


def test_pdf_returns_snapshot_and_matches_selected_sort(authenticated_client, inventory_repository, csrf):
    authenticated_client.post("/one_alma/units/1", data={"sold": "true", "csrf_token": csrf})

    response = authenticated_client.get("/one_alma/pdf?sort=0&descending=true")

    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]
    assert "Продано: 1." in text and "1 343,43" in text


def test_pdf_rejects_invalid_sort_column(authenticated_client, inventory_repository, csrf):
    response = authenticated_client.get("/one_alma/pdf?sort=99")

    assert response.status_code == 422


def test_showcase_endpoint_updates_only_manual_display_count(authenticated_client, inventory_repository, csrf):
    before = inventory_repository.snapshot()

    response = authenticated_client.post("/one_alma/products/1/showcase", data={"quantity": "1", "csrf_token": csrf})
    page = authenticated_client.get("/one_alma")

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.json()["showcase_quantity"] == response.json()["showcased_products"] == 1
    assert response.json()["products"][0]["showcase_quantity"] == 1
    assert "✓ 1 шт." in page.text
    assert inventory_repository.snapshot().valuations == before.valuations
    assert inventory_repository.snapshot().products[0].units == before.products[0].units


def test_sold_count_endpoint_preserves_history_and_returns_updated_snapshot(authenticated_client, inventory_repository, csrf):
    inventory_repository.set_showcase_quantity(1, 2)

    response = authenticated_client.post("/one_alma/products/1/sold", data={"quantity": "2", "csrf_token": csrf})
    page = authenticated_client.get("/one_alma")

    assert response.status_code == 200
    data = response.json()
    assert (data["total"], data["sold"], data["available"], data["showcase_quantity"]) == (50, 2, 48, 0)
    assert len(data["products"]) == 25
    assert all(unit["sold"] for unit in data["products"][0]["units"])
    assert '<span class="count-text">2 из 2</span>' in page.text
    assert data["valuations"][0]["amounts"][1] == "2 686,86"


@pytest.mark.parametrize("kind", ["showcase", "sold"])
@pytest.mark.parametrize("quantity", ["-1", "3", "1.5", "wrong", "1001"])
def test_invalid_product_count_is_rejected_atomically(authenticated_client, inventory_repository, csrf, kind, quantity):
    before = inventory_repository.snapshot()

    response = authenticated_client.post(f"/one_alma/products/1/{kind}", data={"quantity": quantity, "csrf_token": csrf})

    assert response.status_code == 422
    assert inventory_repository.snapshot() == before


@pytest.mark.parametrize("kind", ["showcase", "sold"])
def test_missing_product_count_returns_not_found(authenticated_client, inventory_repository, csrf, kind):
    before = inventory_repository.snapshot()

    response = authenticated_client.post(f"/one_alma/products/9999/{kind}", data={"quantity": "0", "csrf_token": csrf})

    assert response.status_code == 404
    assert inventory_repository.snapshot() == before


def confirm_import(client, csrf, file):
    preview = client.post("/one_alma/import", data={"csrf_token": csrf}, files={"file": file})
    assert preview.status_code == 200
    return client.post(
        "/one_alma/import",
        data={"csrf_token": csrf, "fingerprint": preview.json()["preview"]["fingerprint"]},
        files={"file": file},
    )


def test_excel_preview_does_not_change_inventory_and_shows_replenishment(authenticated_client, inventory_repository, csrf):
    first = inventory_repository.snapshot().products[0]
    row = list(first.cells[1:]) + [3]
    file = excel_file(row, headers=list(IMPORT_HEADERS))
    before = inventory_repository.snapshot()

    response = authenticated_client.post("/one_alma/import", data={"csrf_token": csrf}, files={"file": file})

    preview = response.json()["preview"]
    assert response.status_code == 200 and preview["valid"]
    assert (preview["added"], preview["updated"], preview["quantity"]) == (0, 1, 3)
    assert (preview["rows"][0]["before"], preview["rows"][0]["after"]) == (2, 5)
    assert inventory_repository.snapshot() == before


def test_excel_extra_fields_are_ignored_and_showcase_sales_prices_survive(authenticated_client, inventory_repository, csrf):
    inventory_repository.set_showcase_quantity(1, 1)
    inventory_repository.set_sold_count(1, 1)
    before = inventory_repository.snapshot()
    first = before.products[0]
    row = list(first.cells[1:]) + [3, 99, "=1+1", "Someone"]
    row[4] = 999
    file = excel_file(row, headers=list(IMPORT_HEADERS) + ["Витрина", "Продано", "Кому"])

    response = confirm_import(authenticated_client, csrf, file)
    repeated = confirm_import(authenticated_client, csrf, file)

    inventory = inventory_repository.snapshot()
    assert response.status_code == 201 and repeated.status_code == 200
    assert inventory.products[0].cells == first.cells
    assert inventory.products[0].showcase_quantity == inventory.products[0].sold_count == 1
    assert inventory.products[0].units[:2] == first.units
    assert (len(inventory.products), inventory.total) == (25, 53)


def test_pdf_accepts_showcase_sort_and_contains_display_quantity(authenticated_client, inventory_repository, csrf):
    inventory_repository.set_showcase_quantity(2, 1)

    response = authenticated_client.get("/one_alma/pdf?sort=9&descending=true")

    text = " ".join(" ".join(page.extract_text().split()) for page in PdfReader(BytesIO(response.content)).pages)
    assert response.status_code == 200
    assert "На витрине: 1 шт." in text and "Витрина, шт." in text


def test_excel_upload_repeated_file_preserves_counts_ids_and_sales(authenticated_client, inventory_repository, csrf):
    file = excel_file(["0001", "Крем из Excel", "50 мл", "09.2028", 100, 200, 300, 4], headers=list(IMPORT_HEADERS))
    response = confirm_import(authenticated_client, csrf, file)
    added = inventory_repository.snapshot().products[-1]
    inventory_repository.set_sold(added.units[-1].id, True)
    before = inventory_repository.snapshot()

    repeated = confirm_import(authenticated_client, csrf, file)

    assert response.status_code == 201
    assert response.json()["detail"] == "Добавлено позиций: 1. Пополнено позиций: 0."
    assert before.total == 54 and added.cells[:2] == ("26", "0001")
    assert repeated.status_code == 200
    assert "уже загружена" in repeated.json()["detail"]
    assert InventoryRepository(inventory_repository.database).snapshot() == before


def test_excel_upload_without_numbers_adds_different_expiry_batch(authenticated_client, inventory_repository, csrf):
    first = inventory_repository.snapshot().products[0]
    row = list(first.cells[1:]) + [3]
    different = row.copy()
    different[3] = "10.2029"
    file = excel_file(row, different, headers=list(IMPORT_HEADERS))

    response = confirm_import(authenticated_client, csrf, file)

    inventory = inventory_repository.snapshot()
    assert response.status_code == 201
    assert response.json()["detail"] == "Добавлено позиций: 1. Пополнено позиций: 1."
    assert (len(inventory.products), inventory.total) == (26, 56)
    assert inventory.products[0].cells == first.cells
    assert inventory.products[-1].cells[0] == "26"
    assert inventory.products[-1].cells[4] == "10.2029"


def test_invalid_rows_are_listed_and_cannot_be_partially_imported(authenticated_client, inventory_repository, csrf):
    before = inventory_repository.snapshot()
    file = excel_file(
        ["0001", "Крем", "50 мл", "09.2028", 100, None, None, 4],
        ["0002", "Ошибка количества", "50 мл", "09.2028", 100, None, None, 0],
        ["0003", "Ошибка срока", "50 мл", "99.2028", 100, None, None, 3],
        headers=list(IMPORT_HEADERS),
    )
    preview = authenticated_client.post("/one_alma/import", data={"csrf_token": csrf}, files={"file": file}).json()["preview"]

    response = authenticated_client.post(
        "/one_alma/import", data={"csrf_token": csrf, "fingerprint": preview["fingerprint"]}, files={"file": file}
    )

    assert not preview["valid"]
    assert [row["row"] for row in preview["rows"] if row["error"]] == [3, 4]
    assert response.status_code == 422
    assert "Строка 3" in response.json()["detail"] and "Строка 4" in response.json()["detail"]
    assert inventory_repository.snapshot() == before


def test_stock_overflow_is_highlighted_and_commit_rolls_back(authenticated_client, inventory_repository, csrf):
    before = inventory_repository.snapshot()
    row = list(before.products[0].cells[1:]) + [1000]
    file = excel_file(row, ["NEW", "Новый", "1 мл", "01.2030", None, None, None, 2], headers=list(IMPORT_HEADERS))
    preview = authenticated_client.post("/one_alma/import", data={"csrf_token": csrf}, files={"file": file}).json()["preview"]

    response = authenticated_client.post(
        "/one_alma/import", data={"csrf_token": csrf, "fingerprint": preview["fingerprint"]}, files={"file": file}
    )

    assert not preview["valid"] and "1000" in preview["rows"][0]["error"]
    assert response.status_code == 422
    assert inventory_repository.snapshot() == before


def test_changed_file_after_preview_is_rejected(authenticated_client, inventory_repository, csrf):
    before = inventory_repository.snapshot()
    file = excel_file(["NEW", "Новый", "1 мл", "01.2030", None, None, None, 2], headers=list(IMPORT_HEADERS))

    response = authenticated_client.post(
        "/one_alma/import", data={"csrf_token": csrf, "fingerprint": "0" * 64}, files={"file": file}
    )

    assert response.status_code == 422 and "изменился" in response.json()["detail"]
    assert inventory_repository.snapshot() == before


def test_downloaded_template_can_import_two_identical_distinct_deliveries(authenticated_client, inventory_repository, csrf):
    files = []
    for _ in range(2):
        response = authenticated_client.get("/one_alma/import/template")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        book = load_workbook(BytesIO(response.content))
        book.worksheets[0].append(["NEW", "Новый", "1 мл", "01.2030", None, None, None, 2])
        output = BytesIO()
        book.save(output)
        book.close()
        files.append(("delivery.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))

    responses = [confirm_import(authenticated_client, csrf, file) for file in files]

    assert all(response.status_code == 201 for response in responses)
    assert inventory_repository.snapshot().products[-1].available == 4


def test_manual_add_ignores_submitted_position_and_assigns_next(authenticated_client, inventory_repository, csrf):
    data = product_data(csrf)
    data["position"] = "1"

    response = authenticated_client.post("/one_alma/products", data=data)

    assert response.status_code == 201
    assert inventory_repository.snapshot().products[-1].cells[0] == "26"


@pytest.mark.parametrize(
    "method,path",
    [
        ("GET", "/one_alma"),
        ("GET", "/one_alma/pdf"),
        ("GET", "/one_alma/import/template"),
        ("POST", "/one_alma/units/1"),
        ("POST", "/one_alma/products/1/sold"),
        ("POST", "/one_alma/products/1/showcase"),
        ("POST", "/one_alma/products"),
        ("PUT", "/one_alma/products/1"),
        ("DELETE", "/one_alma/products"),
        ("POST", "/one_alma/import"),
    ],
)
def test_unauthenticated_routes_reject_access_and_do_not_create_db(client, inventory_repository, method, path):
    response = client.request(method, path, follow_redirects=False)

    assert response.status_code in (303, 401, 422)
    assert not inventory_repository.database.exists()
