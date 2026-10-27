import re
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from src.main import application
from src.utils.inventory.excel_import import IMPORT_LABELS, InventoryExcelImporter
from src.utils.inventory.models import InventoryItem
from src.utils.inventory.repository import InventoryRepository
from src.utils.one_alma.excel_import import IMPORT_HEADERS
from src.utils.one_alma.repository import InventoryRepository as OneAlmaRepository


@pytest.fixture
def stores(tmp_path, monkeypatch):
    inventory = InventoryRepository(tmp_path / "inventory.sqlite3")
    one_alma = OneAlmaRepository(tmp_path / "one_alma.sqlite3")
    monkeypatch.setattr(application.routes.inventory, "repository", inventory)
    monkeypatch.setattr(application.routes.one_alma, "repository", one_alma)
    return inventory, one_alma


@pytest.fixture
def tokens(authenticated_client, stores):
    result = {}
    for module in ("inventory", "one_alma"):
        page = authenticated_client.get(f"/{module}")
        token = re.search(r'data-csrf-token="([^"]+)"', page.text)
        assert page.status_code == 200 and token is not None
        result[module] = token[1]
    return result


def excel(headers, rows, leading_rows=0):
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for _ in range(leading_rows):
        sheet.append(["Склад"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return ("stock.xlsx", output.getvalue())


def upload(client, module, token, file, fingerprint=""):
    return client.post(
        f"/{module}/import",
        files={"file": file},
        headers={"X-CSRF-Token": token},
        data={"csrf_token": token, "fingerprint": fingerprint},
    )


def test_both_modules_render_identical_import_dialogs(authenticated_client, stores):
    pages = [authenticated_client.get(f"/{module}").text for module in ("inventory", "one_alma")]

    dialogs = []
    for html in pages:
        matched = re.search(r'<dialog id="import-dialog".*?</dialog>', html, re.S)
        assert matched is not None
        dialogs.append(matched.group(0))

    assert dialogs[0].replace("/inventory/", "/MODULE/") == dialogs[1].replace("/one_alma/", "/MODULE/")
    for html in pages:
        assert "Предварительный просмотр" not in html and "Применить импорт" not in html
        assert ">Проверить файл</button>" in html
        assert html.count("Скачать шаблон Excel") == 1
        assert html.index("/static/import_dialog.js") < html.index(
            "/static/" + ("inventory.js" if "clinic-inventory-page" in html else "one_alma.js")
        )
        assert "/static/css/import_dialog.css" in html


def test_inventory_add_button_is_in_main_actions_and_examples_are_placeholders(authenticated_client, stores):
    html = authenticated_client.get("/inventory").text
    nav_match = re.search(r'<nav class="inventory-actions".*?</nav>', html, re.S)
    form_match = re.search(r'<form id="product-form".*?</form>', html, re.S)
    assert nav_match is not None and form_match is not None
    nav, product_form = nav_match.group(0), form_match.group(0)

    inputs = re.findall(r"<input[^>]+>", product_form)

    assert 'id="add-product"' in nav and html.count('id="add-product"') == 1
    assert "table-toolbar" not in html
    assert all('placeholder="' in field for field in inputs)
    assert all(" value=" not in field for field in inputs if 'name="quantity"' not in field)
    assert 'name="expiry" placeholder="09.2028"' in html


def test_one_alma_keeps_single_row_controls_without_selection_column(authenticated_client, stores):
    html = authenticated_client.get("/one_alma").text

    header_match = re.search(r"<thead><tr>.*?</tr></thead>", html[html.index('<table class="inventory-table"') :], re.S)
    assert header_match is not None
    header = header_match.group(0)

    assert len(re.findall(r"<th\b", header)) == 11
    assert html.count('class="icon-button delete-product"') == 25
    assert html.count('class="icon-button record-sale"') == 25
    assert all(token not in html for token in ("Удалить выбранные", "selection-count", "select-all", "select-product"))


def test_inventory_template_roundtrips_all_supported_fields(authenticated_client, stores):
    response = authenticated_client.get("/inventory/import/template")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.worksheets[0]
    assert sheet is not None
    assert sheet.max_row == 1 and workbook.sheetnames == ["Импорт"]
    headers = [cell.value for cell in sheet[1]]
    sheet.append(["0007", "EMANSI", "Крем", "1200 ₽", "По запросу", datetime(2028, 9, 1), 3])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    imported = InventoryExcelImporter().parse(output.getvalue(), "stock.xlsx")

    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert headers == list(IMPORT_LABELS.values())
    assert imported[0].model_dump() == dict(
        article="0007",
        brand="EMANSI",
        name="Крем",
        purchase_price="1200 ₽",
        sale_price="По запросу",
        expiry="09.2028",
        quantity=3,
    )
    assert stores[0].items() == []


@pytest.mark.parametrize("module", ["inventory", "one_alma"])
def test_bad_rows_are_all_reported_and_commit_cannot_partially_import(authenticated_client, stores, tokens, module):
    if module == "inventory":
        file = excel(
            list(IMPORT_LABELS.values()),
            [
                ["0001", "", "Крем", "", "", "09.2028", 3],
                ["12A", "", "Крем", "", "", "09.2028", 3],
                ["0003", "", "", "", "", "99.2028", 3],
            ],
            leading_rows=2,
        )
    else:
        file = excel(
            list(IMPORT_HEADERS),
            [
                ["NEW1", "Крем", "50 мл", "09.2028", "", "", "", 3],
                ["NEW2", "", "50 мл", "09.2028", "", "", "", 3],
                ["NEW3", "Крем", "50 мл", "99.2028", "", "", "", 3],
            ],
            leading_rows=2,
        )
    before = stores[1].snapshot()
    preview = upload(authenticated_client, module, tokens[module], file)
    result = preview.json()["preview"]

    commit = upload(authenticated_client, module, tokens[module], file, result["fingerprint"])

    assert preview.status_code == 200 and not result["valid"]
    assert [row["row"] for row in result["rows"] if row["error"]] == [5, 6]
    assert "Годен до:" in result["rows"][2]["error"]
    assert commit.status_code == 422 and "Строка 5" in commit.json()["detail"] and "Строка 6" in commit.json()["detail"]
    assert stores[0].items() == [] and stores[1].snapshot() == before


@pytest.mark.parametrize("module", ["inventory", "one_alma"])
def test_same_field_has_same_validation_message_in_both_modules(tokens, authenticated_client, module):
    file = (
        excel(list(IMPORT_LABELS.values()), [["0001", "", "", "", "", "09.2028", 2]])
        if module == "inventory"
        else excel(list(IMPORT_HEADERS), [["0001", "", "50 мл", "09.2028", "", "", "", 2]])
    )

    response = upload(authenticated_client, module, tokens[module], file)

    assert response.status_code == 200
    assert response.json()["preview"]["rows"][0]["error"] == "Наименование: заполните поле"


def test_inventory_duplicate_articles_highlight_every_affected_excel_row(authenticated_client, stores, tokens):
    file = excel(["Артикул", "Наименование", "Количество"], [["0001", "Крем", 2], ["0002", "Гель", 1], ["0001", "Крем", 3]])

    response = upload(authenticated_client, "inventory", tokens["inventory"], file)

    preview = response.json()["preview"]
    assert not preview["valid"]
    assert [row["row"] for row in preview["rows"] if row["error"]] == [2, 4]
    assert "повторяется" in preview["rows"][0]["error"]
    assert not preview["rows"][1]["error"] and stores[0].items() == []


def test_inventory_ambiguous_existing_article_is_a_row_error(authenticated_client, stores, tokens):
    store = stores[0]
    for _ in range(2):
        store.add(InventoryItem(article="0001", name="Крем", quantity=7))
    file = excel(["Артикул", "Наименование", "Количество"], [["0001", "Крем", 3], ["0002", "Гель", 2]])
    before = store.items()
    response = upload(authenticated_client, "inventory", tokens["inventory"], file)
    preview = response.json()["preview"]

    commit = upload(authenticated_client, "inventory", tokens["inventory"], file, preview["fingerprint"])

    assert not preview["valid"] and "нескольких" in preview["rows"][0]["error"]
    assert not preview["rows"][1]["error"]
    assert commit.status_code == 422 and store.items() == before


def test_inventory_replaces_quantity_while_one_alma_adds_delivery(authenticated_client, stores, tokens):
    store, one_alma = stores
    store.add(InventoryItem(article="0001", name="Крем", quantity=7))
    before = one_alma.snapshot()
    product = before.products[0]
    inventory_file = excel(["Артикул", "Наименование", "Количество"], [["0001", "Крем", 3]])
    delivery_file = excel(list(IMPORT_HEADERS), [[*product.cells[1:], 3]])
    inventory_preview = upload(authenticated_client, "inventory", tokens["inventory"], inventory_file).json()["preview"]
    delivery_preview = upload(authenticated_client, "one_alma", tokens["one_alma"], delivery_file).json()["preview"]

    inventory_result = upload(
        authenticated_client, "inventory", tokens["inventory"], inventory_file, inventory_preview["fingerprint"]
    )
    delivery_result = upload(
        authenticated_client, "one_alma", tokens["one_alma"], delivery_file, delivery_preview["fingerprint"]
    )

    assert inventory_result.status_code == delivery_result.status_code == 201
    assert inventory_preview["rows"][0]["before"]["quantity"] == "7"
    assert store.items()[0].quantity == 3
    assert len(one_alma.snapshot().products[0].units) == len(product.units) + 3


@pytest.mark.parametrize("module,valid", [("inventory", True), ("one_alma", False)])
def test_quantity_rules_remain_module_specific(authenticated_client, stores, tokens, module, valid):
    file = (
        excel(["Артикул", "Наименование", "Количество"], [["0001", "Крем", 0]])
        if module == "inventory"
        else excel(list(IMPORT_HEADERS), [["0001", "Крем", "50 мл", "09.2028", "", "", "", 0]])
    )

    response = upload(authenticated_client, module, tokens[module], file)

    assert response.json()["preview"]["valid"] is valid


def test_inventory_template_requires_authentication(client, stores):
    response = client.get("/inventory/import/template", follow_redirects=False)

    assert response.status_code in (303, 401)
    assert stores[0].items() == []


@pytest.mark.parametrize("module", ["inventory", "one_alma"])
def test_row_validation_does_not_bypass_csrf(authenticated_client, stores, tokens, module):
    file = excel(["Артикул", "Наименование", "Количество"], [["0001", "Крем", 3]])

    response = upload(authenticated_client, module, "wrong", file)

    assert response.status_code == 403 and stores[0].items() == []
