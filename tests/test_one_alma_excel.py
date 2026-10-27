from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from src.utils.one_alma.excel_import import ExcelProductImporter

HEADERS = ["№", "Артикул", "Наименование", "Объём", "Годен до", "Цена ЦБ", "Цена опт", "Цена магазина", "Количество"]


def excel_bytes(rows, headers=HEADERS):
    # An in-memory fixture, not a user workbook or deliverable.
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_import_reads_all_columns_and_quantities_and_leading_zero():
    content = excel_bytes([[26, "0001", "Крем", "50 мл", datetime(2028, 9, 1), 100.25, None, "250,10", 3]])

    products = ExcelProductImporter().parse(content, "products.xlsx")

    assert len(products) == 1
    assert products[0].quantity == 3
    assert products[0].to_cells() == ("1", "0001", "Крем", "50 мл", "09.2028", "100,25", "—", "250,10")


def test_quantity_column_is_required_for_deliveries():
    content = excel_bytes([[26, "0001", "Крем", "50 мл", "09.2028", 100, 200, 300]], HEADERS[:-1])

    with pytest.raises(ValueError, match="заголовки"):
        ExcelProductImporter().parse(content, "products.xlsx")


def test_number_is_not_required_and_does_not_affect_fingerprint():
    row = ["0001", "Крем", "50 мл", "09.2028", 100, None, None, 3]
    importer = ExcelProductImporter()

    current = importer.inspect(excel_bytes([row], HEADERS[1:]), "delivery.xlsx")
    legacy = importer.inspect(excel_bytes([[999] + row]), "old.xlsx")

    assert current.valid and current.products == legacy.products
    assert current.fingerprint == legacy.fingerprint


def test_validation_collects_every_bad_row():
    row = ["0001", "Крем", "50 мл", "09.2028", 100, None, None, 3]
    content = excel_bytes(
        [row, ["", "Крем", "50 мл", "09.2028", -1, None, None, 0], row[:3] + ["99.2028"] + row[4:]], HEADERS[1:]
    )

    shipment = ExcelProductImporter().inspect(content, "delivery.xlsx")

    assert not shipment.valid
    assert [row["row"] for row in shipment.rows if row["error"]] == [3, 4]
    assert "Артикул" in str(shipment.rows[1]["error"]) and "Цена ЦБ" in str(shipment.rows[1]["error"])
    assert "Количество" in str(shipment.rows[1]["error"])
    assert "срок годности" in str(shipment.rows[2]["error"])


def test_new_templates_have_distinct_delivery_ids_and_no_number_column():
    first = ExcelProductImporter.template()
    second = ExcelProductImporter.template()
    books = [load_workbook(BytesIO(content)) for content in (first, second)]

    headers = [cell.value for cell in books[0].worksheets[0][1]]

    assert headers == HEADERS[1:]
    assert books[0].worksheets[0].max_row == 1
    assert books[0]["_OneAlma"].sheet_state == "hidden"
    assert books[0]["_OneAlma"]["B1"].value != books[1]["_OneAlma"]["B1"].value
    for book in books:
        book.close()


def test_showcase_columns_are_never_read_from_excel():
    row = [26, "0001", "Крем", "50 мл", "09.2028", 100, 200, 300, 3]
    content = excel_bytes([row + [99, "=1+1", -5]], HEADERS + ["Витрина", "showcase_quantity", "На витрине"])

    products = ExcelProductImporter().parse(content, "products.xlsx")

    assert len(products) == 1 and products[0].quantity == 3
    assert products[0].to_cells() == ("1", "0001", "Крем", "50 мл", "09.2028", "100,00", "200,00", "300,00")


@pytest.mark.parametrize(
    "row",
    [
        [26, "0001", "Крем", "50 мл", "99.2028", 100, 200, 300, 2],
        [26, "0001", "Крем", "50 мл", "09.2028", -10, 200, 300, 2],
        [26, "0001", "Крем", "50 мл", "09.2028", "=1+1", 200, 300, 2],
        [26, "0001", "Крем", "50 мл", "09.2028", 100, 200, 300, 0],
        [26, "0001", "Крем", "50 мл", "09.2028", 100, 200, 300, 1.5],
        [26, "0001", "", "50 мл", "09.2028", 100, 200, 300, 2],
    ],
)
def test_invalid_rows_report_line_and_prevent_import(row):
    content = excel_bytes([row])

    with pytest.raises(ValueError, match="Строка 2"):
        ExcelProductImporter().parse(content, "products.xlsx")


@pytest.mark.parametrize(
    "content, filename, error",
    [
        (b"not excel", "products.xlsx", "Некорректный"),
        (b"", "products.xls", ".xlsx"),
        (excel_bytes([], ["wrong"]), "products.xlsx", "заголовки"),
    ],
)
def test_bad_workbook_and_missing_headers_have_safe_errors(content, filename, error):
    with pytest.raises(ValueError, match=error):
        ExcelProductImporter().parse(content, filename)


def test_malformed_workbook_xml_returns_import_error():
    original = BytesIO(excel_bytes([]))
    corrupted = BytesIO()
    with ZipFile(original) as source, ZipFile(corrupted, "w") as target:
        for item in source.infolist():
            content = b"<broken" if item.filename == "xl/workbook.xml" else source.read(item.filename)
            target.writestr(item, content)

    with pytest.raises(ValueError, match="Не удалось прочитать"):
        ExcelProductImporter().parse(corrupted.getvalue(), "products.xlsx")
